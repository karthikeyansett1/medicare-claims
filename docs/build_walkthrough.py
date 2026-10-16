"""Build the walkthrough DOCX from its reviewed Markdown source.

Run with the Codex bundled Python (python-docx and Pillow); render using the
bundled documents skill render_docx.py, never a desktop LibreOffice install.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'full_project_walkthrough.md'
OUTPUT = ROOT / 'Provider_Performance_Analytics_Full_Walkthrough.docx'

def element(tag, **attrs):
    node=OxmlElement(tag)
    for key,value in attrs.items(): node.set(qn(key),str(value))
    return node

def inline(p, text):
    for bit in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        r=p.add_run(bit[2:-2] if bit.startswith('**') else bit[1:-1] if bit.startswith('`') else bit)
        if bit.startswith('**'): r.bold=True
        if bit.startswith('`'): r.font.name='DejaVu Sans Mono'; r.font.size=Pt(9)

def build():
    text=SOURCE.read_text()
    doc=Document(); sec=doc.sections[0]
    for node in list(doc.styles.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
    sec.top_margin=sec.bottom_margin=Inches(.7)
    sec.left_margin=sec.right_margin=Inches(.8)
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Bullet','List Number']:
        s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(7)
    normal=doc.styles['Normal']; normal.font.size=Pt(11); normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',28),('Heading 1',18),('Heading 2',12),('Heading 3',11)]:
        s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=True
        s.paragraph_format.space_before=Pt(12); s.paragraph_format.keep_with_next=True
    doc.styles['Caption'].font.size=Pt(9); doc.styles['Caption'].font.italic=True
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=foot.add_run(); r.font.size=Pt(9)
    r._r.append(element('w:fldChar',**{'w:fldCharType':'begin'}))
    fld=element('w:instrText'); fld.text='PAGE'; r._r.append(fld)
    r._r.append(element('w:fldChar',**{'w:fldCharType':'end'}))
    lines=text.splitlines(); title=lines[0][2:]
    doc.add_paragraph(title,'Title')
    doc.add_paragraph('Full project walkthrough','Subtitle')
    doc.add_paragraph('Implementation, rationale, verification, and interview preparation','Subtitle')
    doc.add_paragraph('Historical CMS DE SynPUF Samples 1 and 2')
    doc.add_paragraph('October 16, 2026')
    doc.add_paragraph('A practical manual for understanding and reproducing the warehouse, attribution analysis, analytics assistant, evaluation, dashboard, CI, and public deployment.')
    doc.add_page_break(); doc.add_heading('Contents',1)
    headings=[x[3:] for x in lines if x.startswith('## ')]
    for n,h in enumerate(headings):
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
        link=element('w:hyperlink',**{'w:anchor':f'section{n}'})
        run=element('w:r'); props=element('w:rPr'); props.append(element('w:sz',**{'w:val':'20'})); run.append(props)
        t=element('w:t'); t.text=h; run.append(t); link.append(run); p._p.append(link)
    i=1; section=-1; figure=0
    while i<len(lines):
        line=lines[i]
        if not line.strip(): i+=1; continue
        if line.startswith('## '):
            doc.add_page_break(); section+=1
            p=doc.add_heading(line[3:],1)
            p._p.insert(0,element('w:bookmarkStart',**{'w:id':str(section),'w:name':f'section{section}'}))
            p._p.append(element('w:bookmarkEnd',**{'w:id':str(section)})); i+=1; continue
        if line.startswith('### '): doc.add_heading(line[4:],2); i+=1; continue
        if line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line); path=(ROOT/m[2]).resolve()
            w,h=Image.open(path).size; width=min(6.85,5.6*w/h)
            p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pic=p.add_run().add_picture(str(path),width=Inches(width))
            pic._inline.docPr.set('descr',m[1]); figure+=1
            doc.add_paragraph(f'Figure {figure}  {m[1]}','Caption'); i+=1; continue
        if line.startswith('```'):
            i+=1
            while i<len(lines) and not lines[i].startswith('```'):
                p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
                r=p.add_run(lines[i]); r.font.name='DejaVu Sans Mono'; r.font.size=Pt(9); i+=1
            i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                cells=[c.strip() for c in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',c) for c in cells): rows.append(cells)
                i+=1
            tab=doc.add_table(rows=1,cols=len(rows[0])); tab.autofit=False
            widths=([1.6,1.4,1.4,2.45] if len(rows[0])==4 else [3.25,1.8,1.8])
            for col,width in zip(tab.columns,widths): col.width=Inches(width)
            for n,row in enumerate(rows):
                cells=tab.rows[0].cells if n==0 else tab.add_row().cells
                for c,value in enumerate(row):
                    cells[c].width=Inches(widths[c]); inline(cells[c].paragraphs[0],value)
                    tcPr=cells[c]._tc.get_or_add_tcPr();tcPr.append(element('w:vAlign',**{'w:val':'center'}))
                    borders=element('w:tcBorders')
                    for edge in ('top','left','bottom','right'): borders.append(element(f'w:{edge}',**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
                    tcPr.append(borders)
                    margins=element('w:tcMar')
                    for edge in ('top','bottom','left','right'): margins.append(element(f'w:{edge}',**{'w:w':'100','w:type':'dxa'}))
                    tcPr.append(margins)
                    if n==0: tcPr.append(element('w:shd',**{'w:fill':'E7EDF3'}))
                    for p in cells[c].paragraphs:
                        p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                        if c>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:r.font.size=Pt(10);r.bold=n==0
                if n==0: tab.rows[0]._tr.get_or_add_trPr().append(element('w:tblHeader'))
                tab.rows[n]._tr.get_or_add_trPr().append(element('w:cantSplit'))
            doc.add_paragraph();continue
        m=re.match(r'^\d+\. (.*)',line)
        if m: inline(doc.add_paragraph(style='List Number'),m[1]);i+=1;continue
        chunks=[line];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|!\[|```|\||\d+\. )',lines[i]):chunks.append(lines[i]);i+=1
        inline(doc.add_paragraph(),' '.join(chunks))
    doc.core_properties.title=title
    doc.core_properties.subject='Complete technical walkthrough and interview preparation'
    doc.core_properties.author='Medicare Provider Analytics project'
    doc.save(OUTPUT);print(OUTPUT)

if __name__=='__main__':build()
